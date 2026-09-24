import csv
import io
import unittest

from app.knowledge.extractors import (
    CSVExtractor,
    DOCXExtractor,
    ExtractionError,
    PDFExtractor,
    TXTExtractor,
    extract_text,
)


class TestTXTExtractor(unittest.TestCase):
    def test_extracts_plain_text(self):
        text = TXTExtractor().extract("Hello, world.".encode("utf-8"))
        self.assertEqual(text, "Hello, world.")

    def test_rejects_non_utf8_bytes(self):
        with self.assertRaises(ExtractionError):
            TXTExtractor().extract(b"\xff\xfe\x00\x01")

    def test_rejects_empty_file(self):
        with self.assertRaises(ExtractionError):
            TXTExtractor().extract("   \n  ".encode("utf-8"))


class TestCSVExtractor(unittest.TestCase):
    def test_flattens_rows_into_readable_lines(self):
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(["name", "price"])
        writer.writerow(["Widget", "10"])
        writer.writerow(["Gadget", "25"])
        text = CSVExtractor().extract(buf.getvalue().encode("utf-8"))

        self.assertIn("name: Widget", text)
        self.assertIn("price: 10", text)
        self.assertIn("name: Gadget", text)

    def test_skips_fully_blank_rows(self):
        csv_bytes = "name,price\nWidget,10\n,\n".encode("utf-8")
        text = CSVExtractor().extract(csv_bytes)
        self.assertEqual(text.count("name:"), 1)

    def test_rejects_empty_csv(self):
        with self.assertRaises(ExtractionError):
            CSVExtractor().extract(b"")

    def test_rejects_header_only_csv(self):
        with self.assertRaises(ExtractionError):
            CSVExtractor().extract("name,price\n".encode("utf-8"))


class TestPDFExtractor(unittest.TestCase):
    def _make_pdf_bytes(self, text: str) -> bytes:
        from pypdf import PdfWriter

        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        # pypdf has no built-in text-drawing API; embedding real text
        # requires a content stream. For a unit test, exercising the
        # "no extractable text" failure path (a genuinely blank page)
        # is the practically testable case without a PDF-authoring
        # dependency this project doesn't otherwise need.
        buf = io.BytesIO()
        writer.write(buf)
        return buf.getvalue()

    def test_blank_pdf_raises_no_extractable_text(self):
        pdf_bytes = self._make_pdf_bytes("")
        with self.assertRaises(ExtractionError):
            PDFExtractor().extract(pdf_bytes)

    def test_garbage_bytes_raise_extraction_error(self):
        with self.assertRaises(ExtractionError):
            PDFExtractor().extract(b"this is not a pdf at all")


class TestDOCXExtractor(unittest.TestCase):
    def _make_docx_bytes(self, paragraphs: list[str]) -> bytes:
        from docx import Document

        doc = Document()
        for p in paragraphs:
            doc.add_paragraph(p)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def test_extracts_paragraph_text(self):
        docx_bytes = self._make_docx_bytes(["Our return policy is 30 days.", "Contact us for details."])
        text = DOCXExtractor().extract(docx_bytes)
        self.assertIn("Our return policy is 30 days.", text)
        self.assertIn("Contact us for details.", text)

    def test_extracts_table_cell_text(self):
        from docx import Document

        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Product"
        table.rows[0].cells[1].text = "Price"
        buf = io.BytesIO()
        doc.save(buf)

        text = DOCXExtractor().extract(buf.getvalue())
        self.assertIn("Product", text)
        self.assertIn("Price", text)

    def test_rejects_empty_document(self):
        docx_bytes = self._make_docx_bytes([])
        with self.assertRaises(ExtractionError):
            DOCXExtractor().extract(docx_bytes)

    def test_garbage_bytes_raise_extraction_error(self):
        with self.assertRaises(ExtractionError):
            DOCXExtractor().extract(b"this is not a docx at all")


class TestExtractTextDispatch(unittest.TestCase):
    def test_dispatches_by_format(self):
        text = extract_text("txt", "hello".encode("utf-8"))
        self.assertEqual(text, "hello")

    def test_unsupported_format_raises_value_error(self):
        with self.assertRaises(ValueError):
            extract_text("exe", b"whatever")


if __name__ == "__main__":
    unittest.main()
